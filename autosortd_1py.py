import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import requests
import yaml
from watchdog.events import FileSystemEventHandler
from watchdog.observers import Observer


# -----------------------------
# Data models (no pydantic)
# -----------------------------
@dataclass
class LLMDecision:
    doc_type: str
    project: Optional[str] = None
    vendor: Optional[str] = None
    date: Optional[str] = None  # YYYY-MM-DD or None
    suggested_name: str = ""
    confidence: float = 0.0  # 0..1
    reasons: List[str] = field(default_factory=list)
    tags: List[str] = field(default_factory=list)

    @staticmethod
    def from_dict(d: Dict[str, Any]) -> "LLMDecision":
        doc_type = str(d.get("doc_type", "")).strip() or "other"
        project = d.get("project")
        vendor = d.get("vendor")
        date = d.get("date")
        suggested_name = str(d.get("suggested_name", "") or "")
        confidence_raw = d.get("confidence", 0.0)
        try:
            confidence = float(confidence_raw)
        except Exception:
            confidence = 0.0
        confidence = max(0.0, min(1.0, confidence))
        reasons = d.get("reasons") or []
        tags = d.get("tags") or []
        if not isinstance(reasons, list):
            reasons = [str(reasons)]
        if not isinstance(tags, list):
            tags = [str(tags)]
        reasons = [str(x) for x in reasons if str(x).strip()]
        tags = [str(x) for x in tags if str(x).strip()]

        if date is not None:
            date = str(date).strip()
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date):
                date = None

        if project is not None:
            project = str(project).strip() or None
        if vendor is not None:
            vendor = str(vendor).strip() or None

        return LLMDecision(
            doc_type=doc_type,
            project=project,
            vendor=vendor,
            date=date,
            suggested_name=suggested_name,
            confidence=confidence,
            reasons=reasons,
            tags=tags,
        )

    def to_dict(self) -> Dict[str, Any]:
        return {
            "doc_type": self.doc_type,
            "project": self.project,
            "vendor": self.vendor,
            "date": self.date,
            "suggested_name": self.suggested_name,
            "confidence": self.confidence,
            "reasons": self.reasons,
            "tags": self.tags,
        }


# -----------------------------
# Paths
# -----------------------------
@dataclass
class Paths:
    root: Path
    staging: Path
    out: Path
    quarantine: Path
    dup: Path
    logs: Path
    cache: Path
    rules_dir: Path


def ensure_dirs(paths: Paths) -> None:
    for p in [
        paths.root,
        paths.staging,
        paths.out,
        paths.quarantine,
        paths.dup,
        paths.logs,
        paths.cache,
        paths.rules_dir,
    ]:
        p.mkdir(parents=True, exist_ok=True)


# -----------------------------
# IO helpers
# -----------------------------
def atomic_write_json(path: Path, data: Any) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def load_cache(cache_path: Path) -> Dict[str, Any]:
    if not cache_path.exists():
        return {}
    try:
        return json.loads(cache_path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_cache(cache_path: Path, cache: Dict[str, Any]) -> None:
    atomic_write_json(cache_path, cache)


def log_jsonl(log_path: Path, obj: Dict[str, Any]) -> None:
    with log_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")


def sha256_file(p: Path, chunk: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        while True:
            b = f.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def safe_filename(name: str, max_len: int = 140) -> str:
    name = re.sub(r"[\\/:*?\"<>|]+", "_", name).strip()
    name = re.sub(r"\s+", " ", name).strip()
    if len(name) > max_len:
        name = name[:max_len].rstrip()
    return name


def wait_until_stable(p: Path, timeout_s: int = 20) -> bool:
    t0 = time.time()
    last = -1
    stable_hits = 0
    while time.time() - t0 < timeout_s:
        try:
            sz = p.stat().st_size
        except FileNotFoundError:
            return False
        if sz == last and sz > 0:
            stable_hits += 1
        else:
            stable_hits = 0
        if stable_hits >= 2:
            return True
        last = sz
        time.sleep(1)
    return False


# -----------------------------
# Extract snippet (docs only; best-effort)
# -----------------------------
MAX_SNIPPET_CHARS = 2500


def extract_snippet(p: Path) -> str:
    ext = p.suffix.lower()
    try:
        if ext in {".txt", ".md"}:
            return p.read_text(encoding="utf-8", errors="ignore")[:MAX_SNIPPET_CHARS]

        if ext == ".pdf":
            try:
                from pypdf import PdfReader  # optional

                reader = PdfReader(str(p))
                text = ""
                for i in range(min(2, len(reader.pages))):
                    text += (reader.pages[i].extract_text() or "") + "\n"
                return text.strip()[:MAX_SNIPPET_CHARS]
            except Exception:
                return ""

        if ext == ".docx":
            try:
                import docx  # optional

                d = docx.Document(str(p))
                parts: List[str] = []
                for para in d.paragraphs[:30]:
                    t = (para.text or "").strip()
                    if t:
                        parts.append(t)
                return "\n".join(parts)[:MAX_SNIPPET_CHARS]
            except Exception:
                return ""

        if ext in {".xlsx", ".xlsm"}:
            try:
                import openpyxl  # optional

                wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
                ws = wb.worksheets[0]
                rows: List[str] = []
                for r in ws.iter_rows(min_row=1, max_row=20, values_only=True):
                    row = " | ".join([str(x) for x in r if x is not None])
                    row = row.strip()
                    if row:
                        rows.append(row[:200])
                return "\n".join(rows)[:MAX_SNIPPET_CHARS]
            except Exception:
                return ""
    except Exception:
        return ""
    return ""


# -----------------------------
# YAML rule engine
# -----------------------------
@dataclass
class CompiledRule:
    id: str
    type: str
    doc_type: str
    confidence: float
    tags: List[str]
    action: Optional[str]
    reason: Optional[str]
    ext_set: Optional[set] = None
    ext_group: Optional[str] = None
    regex: Optional[re.Pattern] = None
    target: str = "name"


def load_yaml(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except Exception:
        return {}


def compile_rules(
    rules_cfg: Dict[str, Any],
) -> Tuple[Dict[str, List[str]], List[CompiledRule]]:
    ext_groups = rules_cfg.get("ext_groups", {}) or {}
    compiled: List[CompiledRule] = []

    for r in rules_cfg.get("rules", []) or []:
        rid = str(r.get("id", "NO_ID"))
        rtype = str(r.get("type", "")).strip()
        doc_type = str(r.get("doc_type", "other")).strip() or "other"
        conf_raw = r.get("confidence", 0.0)
        try:
            conf = float(conf_raw)
        except Exception:
            conf = 0.0
        conf = max(0.0, min(1.0, conf))

        tags = r.get("tags", []) or []
        if not isinstance(tags, list):
            tags = [str(tags)]
        tags = [str(x) for x in tags if str(x).strip()]
        action = r.get("action")
        reason = r.get("reason")

        cr = CompiledRule(
            id=rid,
            type=rtype,
            doc_type=doc_type,
            confidence=conf,
            tags=tags,
            action=action if isinstance(action, str) else None,
            reason=reason if isinstance(reason, str) else None,
        )

        if rtype == "ext":
            m = r.get("match", []) or []
            if not isinstance(m, list):
                m = [str(m)]
            cr.ext_set = set([str(x).lower() for x in m if str(x).strip()])
        elif rtype == "ext_group":
            cr.ext_group = str(r.get("group", "")).strip() or None
        elif rtype == "regex":
            cr.target = str(r.get("target", "name")).strip() or "name"
            pat = str(r.get("pattern", ".*"))
            try:
                cr.regex = re.compile(pat, re.I)
            except Exception:
                continue
        else:
            continue

        compiled.append(cr)

    return ext_groups, compiled


def first_match_rule(
    filename: str,
    fullpath: str,
    ext: str,
    ext_groups: Dict[str, List[str]],
    compiled_rules: List[CompiledRule],
) -> Optional[LLMDecision]:
    for r in compiled_rules:
        ok = False
        if r.type == "ext" and r.ext_set is not None:
            ok = ext in r.ext_set
        elif r.type == "ext_group" and r.ext_group:
            group_exts = [
                str(x).lower() for x in (ext_groups.get(r.ext_group, []) or [])
            ]
            ok = ext in set(group_exts)
        elif r.type == "regex" and r.regex:
            target_str = filename if r.target == "name" else fullpath
            ok = bool(r.regex.search(target_str))

        if ok:
            return LLMDecision(
                doc_type=r.doc_type,
                suggested_name=filename,
                confidence=r.confidence,
                reasons=[f"rule:{r.id}"],
                tags=list(r.tags),
            )
    return None


# -----------------------------
# Mapping (doc_type_map + tag_overrides + gate)
# -----------------------------
@dataclass
class RoutingResult:
    dest_dir: Path
    rename_policy: str  # keep | normalize
    forced_quarantine: bool
    forced_reason: str


def split_rel_path(rel: str) -> Path:
    rel = rel.replace("/", "\\")
    parts = [p for p in rel.split("\\") if p]
    return Path(*parts)


def resolve_destination(
    paths: Paths, mapping_cfg: Dict[str, Any], decision: LLMDecision
) -> RoutingResult:
    doc_type_map = mapping_cfg.get("doc_type_map", {}) or {}
    tag_overrides = mapping_cfg.get("tag_overrides", []) or []
    rename_policy_map = mapping_cfg.get("rename_policy", {}) or {}
    apply_gate = mapping_cfg.get("apply_gate", {}) or {}

    forced_quarantine = False
    forced_reason = ""

    quarantine_doc_types = set(apply_gate.get("quarantine_doc_types", []) or [])
    if decision.doc_type in quarantine_doc_types:
        forced_quarantine = True
        forced_reason = "forced_quarantine_doc_type"

    base_rel = doc_type_map.get(
        decision.doc_type, doc_type_map.get("other", r"Docs\Other")
    )
    dest_rel = base_rel
    rp = rename_policy_map.get(decision.doc_type, "keep")

    tags_set = set([t for t in (decision.tags or [])])
    for ov in tag_overrides:
        need = set(ov.get("when_all", []) or [])
        if need and need.issubset(tags_set):
            dest_rel = ov.get("dest_rel", dest_rel)
            rp = ov.get("rename_policy", rp)
            break

    dest_dir = paths.out / split_rel_path(str(dest_rel))
    rp = str(rp or "keep").strip().lower()
    if rp not in {"keep", "normalize"}:
        rp = "keep"

    return RoutingResult(
        dest_dir=dest_dir,
        rename_policy=rp,
        forced_quarantine=forced_quarantine,
        forced_reason=forced_reason,
    )


def should_auto_apply(
    mapping_cfg: Dict[str, Any], decision: LLMDecision
) -> Tuple[bool, str]:
    apply_gate = mapping_cfg.get("apply_gate", {}) or {}
    q_below = float(apply_gate.get("quarantine_below", 0.90))
    allow = set(apply_gate.get("allow_auto_apply_doc_types", []) or [])

    if decision.confidence < q_below:
        return (False, "low_confidence")
    if allow and decision.doc_type not in allow:
        return (False, "not_in_allowlist_doc_type")
    return (True, "auto_apply")


# -----------------------------
# LLM client (OpenAI-compatible)
# -----------------------------
def _extract_json_obj(text: str) -> Optional[str]:
    s = text.strip()
    if not s:
        return None
    if s.startswith("{") and s.endswith("}"):
        return s
    i = s.find("{")
    j = s.rfind("}")
    if i >= 0 and j > i:
        return s[i : j + 1]
    return None


def llm_classify(
    base_url: str, file_meta: Dict[str, Any], snippet: str, timeout_s: int = 120
) -> Optional[LLMDecision]:
    system = (
        "Return ONLY one JSON object (no markdown, no extra text). "
        "Keys: doc_type, project, vendor, date, suggested_name, confidence, reasons, tags. "
        "doc_type must be one of: ops_doc, other, dev_archive, dev_note, photo. "
        "date must be YYYY-MM-DD or null. confidence is 0..1. "
        "reasons/tags are arrays of short strings."
    )

    user_obj = {
        "file_meta": file_meta,
        "content_snippet": snippet[:MAX_SNIPPET_CHARS],
        "naming_rule": "suggested_name: keep short; prefer original filename unless clear.",
    }

    payload = {
        "model": "not-used",
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": json.dumps(user_obj, ensure_ascii=False)},
        ],
        "temperature": 0,
        "max_tokens": 400,
        "stream": False,
    }

    r = requests.post(
        f"{base_url.rstrip('/')}/chat/completions", json=payload, timeout=timeout_s
    )
    r.raise_for_status()
    content = r.json()["choices"][0]["message"]["content"]
    js = _extract_json_obj(content)
    if js is None:
        return None
    try:
        data = json.loads(js)
    except Exception:
        return None
    dec = LLMDecision.from_dict(data if isinstance(data, dict) else {})
    if not dec.suggested_name:
        dec.suggested_name = str(file_meta.get("name", "unnamed"))
    return dec


def warmup(llm_base_url: str) -> None:
    payload = {
        "model": "not-used",
        "messages": [{"role": "user", "content": "ping"}],
        "max_tokens": 1,
        "temperature": 0,
        "stream": False,
    }
    try:
        requests.post(
            f"{llm_base_url.rstrip('/')}/chat/completions", json=payload, timeout=20
        ).raise_for_status()
    except Exception:
        pass


# -----------------------------
# File move + ledger
# -----------------------------
def move_with_ledger(
    paths: Paths,
    src: Path,
    dst_dir: Path,
    new_name: str,
    ledger_path: Path,
    run_id: str,
    sha: str,
    reason: str,
) -> Path:
    dst_dir.mkdir(parents=True, exist_ok=True)

    new_name = safe_filename(new_name)
    if not new_name:
        new_name = safe_filename(src.stem)

    if not new_name.lower().endswith(src.suffix.lower()):
        new_name = f"{new_name}{src.suffix}"

    dst = dst_dir / new_name

    if dst.exists():
        base = dst.stem
        for i in range(1, 1000):
            cand = dst_dir / f"{base}__{i:03d}{dst.suffix}"
            if not cand.exists():
                dst = cand
                break

    staging_target = paths.staging / f"{run_id}__{src.name}"

    shutil.move(str(src), str(staging_target))
    shutil.move(str(staging_target), str(dst))

    log_jsonl(
        ledger_path,
        {
            "ts": datetime.utcnow().isoformat(timespec="seconds") + "Z",
            "run_id": run_id,
            "action": "move",
            "sha256": sha,
            "reason": reason,
            "before": str(src),
            "after": str(dst),
        },
    )
    return dst


# -----------------------------
# Core handler
# -----------------------------
def handle_file(
    paths: Paths,
    llm_base_url: str,
    rules_cfg: Dict[str, Any],
    mapping_cfg: Dict[str, Any],
    ext_groups: Dict[str, List[str]],
    compiled_rules: List[CompiledRule],
    cache_path: Path,
    ledger_path: Path,
    p: Path,
) -> None:
    if not p.exists() or p.is_dir():
        return

    filename = p.name
    fullpath = str(p)
    ext = p.suffix.lower()

    ignore_exts = set(
        [
            str(x).lower()
            for x in (rules_cfg.get("ignore", {}).get("extensions", []) or [])
        ]
    )
    ignore_globs = rules_cfg.get("ignore", {}).get("globs", []) or []
    if ext in ignore_exts:
        return
    for g in ignore_globs:
        if Path(filename).match(g):
            return

    st_cfg = rules_cfg.get("stability_check", {}) or {}
    st_ignore_exts = set(
        [str(x).lower() for x in (st_cfg.get("ignore_extensions", []) or [])]
    )
    if ext in st_ignore_exts:
        return

    if st_cfg.get("enabled", True):
        timeout_s = int(st_cfg.get("timeout_seconds", 20))
        if not wait_until_stable(p, timeout_s=timeout_s):
            run_id = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
            sha = "unreadable"
            move_with_ledger(
                paths,
                p,
                paths.quarantine,
                f"UNSTABLE__{p.stem}",
                ledger_path,
                run_id,
                sha,
                "unstable_file",
            )
            return

    run_id = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    sha = sha256_file(p)

    cache = load_cache(cache_path)
    if sha in cache:
        move_with_ledger(
            paths,
            p,
            paths.dup,
            f"DUP__{p.stem}",
            ledger_path,
            run_id,
            sha,
            "duplicate_hash",
        )
        return

    decision = first_match_rule(filename, fullpath, ext, ext_groups, compiled_rules)

    docs_exts = set([str(x).lower() for x in (ext_groups.get("docs", []) or [])])

    if decision is None or decision.confidence < 0.92:
        if ext in docs_exts:
            snippet = extract_snippet(p)
            meta = {"name": filename, "ext": ext, "size": p.stat().st_size}
            try:
                llm_dec = llm_classify(llm_base_url, meta, snippet)
            except requests.Timeout:
                move_with_ledger(
                    paths,
                    p,
                    paths.quarantine,
                    f"LLMTIMEOUT__{p.stem}",
                    ledger_path,
                    run_id,
                    sha,
                    "llm_timeout",
                )
                return
            except requests.RequestException:
                move_with_ledger(
                    paths,
                    p,
                    paths.quarantine,
                    f"LLMERROR__{p.stem}",
                    ledger_path,
                    run_id,
                    sha,
                    "llm_error",
                )
                return
            if llm_dec is None:
                move_with_ledger(
                    paths,
                    p,
                    paths.quarantine,
                    f"PARSEFAIL__{p.stem}",
                    ledger_path,
                    run_id,
                    sha,
                    "llm_parse_fail",
                )
                return
            if decision is not None and decision.tags and not llm_dec.tags:
                llm_dec.tags = decision.tags
            decision = llm_dec
        else:
            decision = LLMDecision(
                doc_type="other",
                suggested_name=filename,
                confidence=0.0,
                reasons=["no_rule_non_doc"],
            )

    routing = resolve_destination(paths, mapping_cfg, decision)

    if decision.doc_type.startswith("dev_") or decision.doc_type in {
        "dev_code",
        "dev_repo",
        "dev_config",
    }:
        rename_policy = "keep"
    else:
        rename_policy = routing.rename_policy or "keep"

    if routing.forced_quarantine:
        dst_dir = paths.quarantine
        reason = routing.forced_reason or "forced_quarantine"
    else:
        ok, gate_reason = should_auto_apply(mapping_cfg, decision)
        if ok:
            dst_dir = routing.dest_dir
            reason = gate_reason
        else:
            dst_dir = paths.quarantine
            reason = gate_reason

    if rename_policy == "keep":
        new_name = filename
    else:
        new_name = decision.suggested_name or filename

    final = move_with_ledger(
        paths, p, dst_dir, new_name, ledger_path, run_id, sha, reason
    )

    cache[sha] = {
        "run_id": run_id,
        "after": str(final),
        "decision": decision.to_dict(),
        "reason": reason,
        "dest_dir": str(dst_dir),
    }
    save_cache(cache_path, cache)


# -----------------------------
# Watchdog handler
# -----------------------------
class Handler(FileSystemEventHandler):
    def __init__(
        self,
        paths: Paths,
        llm_base_url: str,
        rules_cfg: Dict[str, Any],
        mapping_cfg: Dict[str, Any],
        ext_groups: Dict[str, List[str]],
        compiled_rules: List[CompiledRule],
        cache_path: Path,
        ledger_path: Path,
    ):
        self.paths = paths
        self.llm = llm_base_url
        self.rules_cfg = rules_cfg
        self.mapping_cfg = mapping_cfg
        self.ext_groups = ext_groups
        self.compiled_rules = compiled_rules
        self.cache_path = cache_path
        self.ledger_path = ledger_path

    def _handle(self, p: Path):
        try:
            handle_file(
                self.paths,
                self.llm,
                self.rules_cfg,
                self.mapping_cfg,
                self.ext_groups,
                self.compiled_rules,
                self.cache_path,
                self.ledger_path,
                p,
            )
        except Exception:
            try:
                if p.exists() and not p.is_dir():
                    run_id = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
                    sha = sha256_file(p)
                    move_with_ledger(
                        self.paths,
                        p,
                        self.paths.quarantine,
                        f"ERROR__{p.stem}",
                        self.ledger_path,
                        run_id,
                        sha,
                        "exception",
                    )
            except Exception:
                pass

    def on_created(self, event):
        if not event.is_directory:
            self._handle(Path(event.src_path))

    def on_moved(self, event):
        if not event.is_directory:
            self._handle(Path(event.dest_path))


def sweep_existing(watch_dir: Path, handler: Handler) -> None:
    for p in watch_dir.iterdir():
        if p.is_file():
            handler._handle(p)


# -----------------------------
# Main
# -----------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=r"C:\_AUTOSORT")
    ap.add_argument("--watch", default=str(Path.home() / "Downloads"))
    ap.add_argument("--llm", default="http://127.0.0.1:8080/v1")
    ap.add_argument("--rules_dir", default=r"C:\_AUTOSORT\rules")
    ap.add_argument(
        "--sweep", action="store_true", help="process existing files at startup"
    )
    args = ap.parse_args()

    root = Path(args.root)
    paths = Paths(
        root=root,
        staging=root / "staging",
        out=root / "out",
        quarantine=root / "quarantine",
        dup=root / "dup",
        logs=root / "logs",
        cache=root / "cache",
        rules_dir=Path(args.rules_dir),
    )
    ensure_dirs(paths)

    rules_path = paths.rules_dir / "rules.yaml"
    mapping_path = paths.rules_dir / "mapping.yaml"

    rules_cfg = load_yaml(rules_path)
    mapping_cfg = load_yaml(mapping_path)

    if not rules_cfg or not rules_cfg.get("rules"):
        print(
            f"Error: rules.yaml not found or empty in {paths.rules_dir}. "
            "Copy project rules/ to that path or set --rules_dir to project rules folder.",
            file=sys.stderr,
        )
        raise SystemExit(1)

    ext_groups, compiled_rules = compile_rules(rules_cfg)
    print(f"Rules loaded: {len(compiled_rules)} rules from {rules_path}")

    cache_path = paths.cache / "cache.json"
    ledger_path = paths.logs / "ledger.jsonl"

    warmup(args.llm)
    print(f"LLM: {args.llm}")

    watch_dir = Path(args.watch)
    if not watch_dir.is_dir():
        print(f"Error: --watch path is not a directory: {watch_dir}", file=sys.stderr)
        raise SystemExit(1)
    print(f"Watching: {watch_dir}")

    observer = Observer()
    h = Handler(
        paths,
        args.llm,
        rules_cfg,
        mapping_cfg,
        ext_groups,
        compiled_rules,
        cache_path,
        ledger_path,
    )
    observer.schedule(h, str(watch_dir), recursive=False)
    observer.start()

    if args.sweep:
        sweep_existing(watch_dir, h)

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
    observer.join()


if __name__ == "__main__":
    main()
